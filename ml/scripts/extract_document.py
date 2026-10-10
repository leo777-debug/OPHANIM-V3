"""Offline extraction. Customer text is never sent to a model or training path."""
import base64
import csv
import io
import json
import os
import sys
import zipfile
from email import policy
from email.parser import BytesParser
from pathlib import Path

MAX_TEXT = 8_000_000


def check_archive(data):
    with zipfile.ZipFile(io.BytesIO(data)) as archive:
        if sum(i.file_size for i in archive.infolist()) > 64 * 1024**2 or any(i.file_size > 32 * 1024**2 for i in archive.infolist()):
            raise ValueError('Office archive exceeds extraction limits')


def ocr(image):
    import pytesseract
    if os.environ.get('OPHANIM_TESSERACT_CMD'):
        pytesseract.pytesseract.tesseract_cmd = os.environ['OPHANIM_TESSERACT_CMD']
    available = pytesseract.get_languages(config='')
    if not {'eng', 'ara'}.issubset(available):
        raise RuntimeError('OCR requires Tesseract with both eng and ara language packs')
    return pytesseract.image_to_string(image, lang='eng+ara', timeout=45)


def extract(name, data):
    ext = Path(name).suffix.lower()
    blocks = []
    def add(text, reference=None, table=False):
        if text.strip():
            blocks.append(dict(text=text.strip(), reference=reference or {}, table=table))
        if sum(len(b['text']) for b in blocks) > MAX_TEXT or len(blocks) > 10000:
            raise ValueError('Extracted text exceeds limit')
    if ext in ('.txt', '.md', '.json', '.jsonl'):
        heading = None
        for paragraph in data.decode('utf-8-sig').split('\n\n'):
            if paragraph.lstrip().startswith('#'):
                heading = paragraph.splitlines()[0].lstrip('# ').strip()
            add(paragraph, {'heading': heading} if heading else {})
    elif ext == '.csv':
        rows = csv.reader(io.StringIO(data.decode('utf-8-sig')))
        header = next(rows, [])
        for i, row in enumerate(rows, 2):
            add(' | '.join(header) + '\n' + ' | '.join(row), {'rowStart': i, 'rowEnd': i}, True)
    elif ext == '.xlsx':
        check_archive(data)
        from openpyxl import load_workbook
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        try:
            for sheet in workbook:
                header = None
                for i, row in enumerate(sheet.iter_rows(values_only=True), 1):
                    values = [str(v) if v is not None else '' for v in row]
                    if not any(values):
                        continue
                    if header is None:
                        header = ' | '.join(values)
                    add(header + '\n' + ' | '.join(values), {'sheet': sheet.title, 'rowStart': i, 'rowEnd': i}, True)
        finally:
            workbook.close()
    elif ext == '.docx':
        check_archive(data)
        from docx import Document
        doc = Document(io.BytesIO(data))
        heading = None
        # iter_inner_content preserves paragraph/table order.
        from docx.text.paragraph import Paragraph
        for item in doc.iter_inner_content():
            if isinstance(item, Paragraph):
                if item.style.name.startswith('Heading'):
                    heading = item.text
                add(item.text, {'heading': heading} if heading else {})
            else:
                add('\n'.join(' | '.join(cell.text for cell in row.cells) for row in item.rows), {'heading': heading} if heading else {}, True)
    elif ext == '.pdf':
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        if len(reader.pages) > 200:
            raise ValueError('PDF exceeds 200-page limit')
        for i, page in enumerate(reader.pages):
            text = page.extract_text(extraction_mode='layout') or ''
            if len(text.strip()) < 30:
                import pypdfium2 as pdfium
                with pdfium.PdfDocument(data) as pdf:
                    bitmap = pdf[i].render(scale=2)
                    text = ocr(bitmap.to_pil())
            add(text, {'page': i + 1})
    elif ext in ('.png', '.jpg', '.jpeg'):
        from PIL import Image
        image = Image.open(io.BytesIO(data))
        image.verify()
        add(ocr(Image.open(io.BytesIO(data))), {'page': 1})
    elif ext == '.eml':
        message = BytesParser(policy=policy.default).parsebytes(data)
        add('\n'.join(f'{k}: {message.get(k, "")}' for k in ('Subject', 'Date', 'From', 'To')), {'heading': 'Email headers'})
        for part in message.walk():
            if part.get_content_type() == 'text/plain' and part.get_content_disposition() != 'attachment':
                add(str(part.get_content()), {'heading': 'Email body'})
        if not any(b['reference'].get('heading') == 'Email body' for b in blocks):
            raise ValueError('Email has no plain-text body; export a plain-text copy')
    else:
        raise ValueError('Unsupported format; .msg is excluded (no reviewed permissive parser)')
    if not blocks:
        raise ValueError('No readable text found')
    return blocks


if __name__ == '__main__':
    # Windows pipe defaults can be CP1252. Node sends/reads UTF-8 JSON,
    # including Arabic filenames and extracted text; use an explicit protocol.
    for stream in (sys.stdin, sys.stdout, sys.stderr):
        stream.reconfigure(encoding='utf-8')
    try:
        request = json.load(sys.stdin)
        raw = base64.b64decode(request['content_base64'], validate=True)
        if len(raw) > 10 * 1024**2:
            raise ValueError('File exceeds 10 MB')
        print(json.dumps({'blocks': extract(request['file_name'], raw)}, ensure_ascii=False))
    except Exception as error:
        print(str(error), file=sys.stderr)
        sys.exit(1)
