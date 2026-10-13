import io
import json
import sys
import tempfile
import unittest
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'scripts'))
sys.path.insert(0,str(Path(__file__).resolve().parents[1]/'eval'))
from extract_document import extract
from generate_dataset import generate,container_number
from model_store import root,inside,register,verify,reserve_destination
from compare import score
from import_knowledge import normalise,nearest,archive,utc
from generate_pretraining_corpus import generate as generate_corpus
from pretrain import load_corpus,pack,verified_base


class PipelineTests(unittest.TestCase):
    def test_download_budget_fits_kaggle_volume_and_still_rejects_insufficient_space(self):
        import os
        from unittest.mock import patch
        from collections import namedtuple
        from download_base_models import download
        usage=namedtuple('usage','total used free')
        with tempfile.TemporaryDirectory(dir=root()/'eval') as directory:
            with patch.dict(os.environ,{'OPHANIM_MODELS_DIR':directory}):
                info={'sha':'fixture-revision','siblings':[{'rfilename':'model.safetensors','size':8*1024**3}]}
                response=io.BytesIO(json.dumps(info).encode())
                def fake_download(repo,name,revision,local_dir):
                    file=Path(local_dir)/name;file.write_bytes(b'synthetic download fixture');return str(file)
                # 11 GiB is less than the old fixed 12/20 GiB check, but fits
                # 8 GiB base + 2 GiB adapter checkpoints + 512 MiB reserve.
                with patch('shutil.disk_usage',return_value=usage(20*1024**3,9*1024**3,11*1024**3)),patch('urllib.request.urlopen',return_value=response),patch('huggingface_hub.hf_hub_download',side_effect=fake_download) as fetch:
                    download('fixture/model','base',extra_required_bytes=2*1024**3)
                    fetch.assert_called_once()
                response=io.BytesIO(json.dumps({'sha':'fixture','siblings':[{'rfilename':'other.safetensors','size':16*1024**3}]}).encode())
                with patch('shutil.disk_usage',return_value=usage(20*1024**3,9*1024**3,11*1024**3)),patch('urllib.request.urlopen',return_value=response),patch('huggingface_hub.hf_hub_download') as fetch:
                    with self.assertRaisesRegex(RuntimeError,'Insufficient free space'):
                        download('fixture/too-large','base',extra_required_bytes=2*1024**3)
                    fetch.assert_not_called()

    def test_kaggle_accepts_expanded_inputs_and_rejects_corruption_and_traversal(self):
        import hashlib,zipfile
        from kaggle_inputs import stage_inputs
        with tempfile.TemporaryDirectory(dir=root()/'eval') as directory:
            folder=Path(directory);inputs=folder/'input';inputs.mkdir()
            checks={'corpus/train.jsonl':hashlib.sha256(b'synthetic fixture').hexdigest()}
            archive=inputs/'ophanim-cpt-input.zip'
            with zipfile.ZipFile(archive,'w') as z:
                z.writestr('corpus/train.jsonl',b'synthetic fixture');z.writestr('bundle_manifest.json',json.dumps(checks))
            self.assertEqual(stage_inputs(inputs,folder/'zip-output')['form'],'zip')
            with self.assertRaisesRegex(RuntimeError,'pinned submission bundle'):
                stage_inputs(inputs,folder/'wrong-version','0'*64)
            # Explicitly simulate Kaggle's observed automatic expansion.
            expanded=folder/'expanded';(expanded/'corpus').mkdir(parents=True)
            (expanded/'corpus/train.jsonl').write_bytes(b'synthetic fixture')
            (expanded/'bundle_manifest.json').write_text(json.dumps(checks))
            self.assertEqual(stage_inputs(expanded,folder/'expanded-output')['verified_files'],1)
            (expanded/'corpus/train.jsonl').write_bytes(b'corrupt')
            with self.assertRaisesRegex(RuntimeError,'Corrupted input'):stage_inputs(expanded,folder/'bad-output')
            (expanded/'bundle_manifest.json').write_text(json.dumps({'../outside':checks['corpus/train.jsonl']}))
            with self.assertRaisesRegex(RuntimeError,'Unsafe input path'):stage_inputs(expanded,folder/'bad-output')

    def test_model_loading_rejects_unregistered_inputs_and_mixed_revisions(self):
        import os
        from unittest.mock import patch
        with tempfile.TemporaryDirectory(dir=root()/'eval') as directory:
            with patch.dict(os.environ,{'OPHANIM_MODELS_DIR':directory}):
                store=root();folder=store/'base'/'synthetic-fixture';folder.mkdir()
                weight=folder/'fixture.safetensors';weight.write_bytes(b'synthetic checksum fixture; not model weights')
                register(store,weight,'synthetic-fixture','v1','none','MIT')
                self.assertEqual(verified_base(store,'base/synthetic-fixture','synthetic-fixture'),(folder,'v1'))
                config=folder/'config.json';config.write_text('{}')
                with self.assertRaises(RuntimeError):verified_base(store,'base/synthetic-fixture','synthetic-fixture')
                register(store,config,'different-model','v1','none','MIT')
                with self.assertRaisesRegex(RuntimeError,'provenance'):verified_base(store,'base/synthetic-fixture','synthetic-fixture')
                register(store,config,'synthetic-fixture','v2','none','MIT')
                with self.assertRaisesRegex(RuntimeError,'Mixed base revisions'):verified_base(store,'base/synthetic-fixture','synthetic-fixture')

    def test_cpu_frozen_linear_preserves_input_gradients_and_adapter_training(self):
        import torch
        from cpu_training import optimise_frozen_linears
        torch.manual_seed(42)
        layer=torch.nn.Linear(8,12).half()
        for p in layer.parameters():p.requires_grad_(False)
        adapter=torch.nn.Linear(12,2)
        x=torch.randn(3,8,dtype=torch.float16,requires_grad=True)
        gold=x.detach().clone().requires_grad_(True)
        target=torch.nn.functional.linear(gold.float(),layer.weight.float(),layer.bias.float()).half()
        target.sum().backward()
        self.assertEqual(optimise_frozen_linears(layer),1)
        actual=layer(x);actual.sum().backward()
        self.assertTrue(torch.equal(actual,target));self.assertTrue(torch.equal(x.grad,gold.grad))
        self.assertEqual(optimise_frozen_linears(adapter),0)
        adapter(actual.detach().float()).sum().backward()
        self.assertIsNotNone(adapter.weight.grad)

    def test_pretraining_uses_unlabelled_text_and_rejects_leakage_or_customer_data(self):
        with tempfile.TemporaryDirectory(dir=root()/'eval') as directory:
            folder=Path(directory);summary=generate_corpus(folder,2,1)
            self.assertEqual(summary['splits']['train']['documents'],12)
            corpus=load_corpus(folder)
            self.assertTrue(all('messages' not in r and 'target' not in r for r in corpus['train']))
            class Tokenizer:
                eos_token_id=999
                def __call__(self,text,**kwargs):return {'input_ids':[ord(c) for c in text]}
            blocks=pack([{'text':'abcd'},{'text':'ef'}],Tokenizer(),4)
            self.assertEqual(blocks,[{'input_ids':[97,98,99,100]},{'input_ids':[999,101,102,999]}])
            file=folder/'train.jsonl';rows=corpus['train'];rows[0]['customer_data']=True
            file.write_text('\n'.join(json.dumps(r) for r in rows),encoding='utf-8')
            with self.assertRaisesRegex(RuntimeError,'customer uploads'):load_corpus(folder)
            rows[0]['customer_data']=False;rows.append(dict(corpus['validation'][0],split='train'))
            file.write_text('\n'.join(json.dumps(r) for r in rows),encoding='utf-8')
            with self.assertRaises(RuntimeError):load_corpus(folder)

    def test_shared_normalisation_and_nearest_reference_points(self):
        source={'features':[{'id':'fixture','geometry':{'coordinates':[55.3,25.2,10]},'properties':{'place':'Synthetic fixture','time':1609459200000,'mag':4.2,'url':'https://example.test/event'}}]}
        record=list(normalise('usgs',source))[0]
        self.assertEqual(record['start'],'2021-01-01T00:00:00+00:00')
        self.assertEqual(record['place']['longitude'],55.3)
        nearest(record,[{'type':'airport','code':'SYN','latitude':25.2,'longitude':55.3}])
        self.assertEqual(record['place']['nearest_airports'][0],{'code':'SYN','distance_km':0.0})
        self.assertEqual(utc('2021-01-01T04:00:00+04:00').hour,0)

    def test_shared_archive_preserves_source_dates(self):
        with tempfile.TemporaryDirectory(dir=root()/'eval') as directory:
            folder=Path(directory);(folder/'past').mkdir();(folder/'present').mkdir()
            file=folder/'present'/'fixture.json';file.write_text('{"events":[]}')
            metadata={'published_at':'2021-01-01T00:00:00Z','valid_from':'2020-12-31','knowledge_period':'present'}
            Path(str(file)+'.meta.json').write_text(json.dumps(metadata))
            archive(folder)
            moved=json.loads((folder/'past'/'fixture.json.meta.json').read_text())
            self.assertEqual(moved['published_at'],metadata['published_at']);self.assertEqual(moved['valid_from'],metadata['valid_from']);self.assertEqual(moved['knowledge_period'],'past')

    def test_synthetic_splits_and_unanswerables(self):
        with tempfile.TemporaryDirectory(dir=root()/'eval') as directory:
            folder=Path(directory);summary=generate(folder,132)
            self.assertEqual(summary['total'],744)
            hashes=[]
            for split in ('train','validation','test'):
                rows=[json.loads(l) for l in (folder/f'{split}.jsonl').read_text(encoding='utf-8').splitlines()]
                self.assertEqual(sum(r['mode']=='general' for r in rows),sum(r['mode']=='logistics' for r in rows))
                qa=[r for r in rows if r['task']=='grounded_qa']
                self.assertGreaterEqual(sum(r['unanswerable'] for r in qa)/len(qa),0.2)
                self.assertTrue(all(r['customer_data'] is False for r in rows))
                from jsonschema import validate
                for row in rows:validate(row['target'],row['schema'])
                hashes.append({json.dumps(r['messages'][:2],sort_keys=True) for r in rows})
            self.assertFalse(hashes[0]&hashes[1]);self.assertFalse(hashes[0]&hashes[2]);self.assertFalse(hashes[1]&hashes[2])

    def test_container_check_digit_known_algorithm(self):
        letters=dict(zip('ABCDEFGHIJKLMNOPQRSTUVWXYZ',[10,12,13,14,15,16,17,18,19,20,21,23,24,25,26,27,28,29,30,31,32,34,35,36,37,38]))
        for serial in (1,42,202600):
            number=container_number(serial);value=sum((int(c) if c.isdigit() else letters[c])*2**i for i,c in enumerate(number[:-1]))
            self.assertEqual(int(number[-1]),value%11%10)

    def test_multilingual_text_csv_and_eml(self):
        self.assertIn('الشحنة',extract('source.md','# Cargo\n\nالشحنة 48'.encode())[1]['text'])
        csv=extract('cargo.csv',b'kind,pieces\nboxes,48\n');self.assertEqual(csv[0]['reference']['rowStart'],2)
        email=extract('alert.eml',b'Subject: Delay\nContent-Type: text/plain; charset=utf-8\n\nShipment delayed.');self.assertIn('delayed',email[-1]['text'])
        with self.assertRaises(ValueError):extract('unsafe.msg',b'unknown')

    def test_extractor_cli_uses_utf8_even_with_windows_cp1252_environment(self):
        import base64,os,subprocess
        request={'file_name':'الشحنة.txt','content_base64':base64.b64encode('الشحنة تحتوي على ٤٨ صندوقاً'.encode()).decode()}
        result=subprocess.run([sys.executable,str(Path(__file__).resolve().parents[1]/'scripts'/'extract_document.py')],
                              input=json.dumps(request,ensure_ascii=False).encode(),capture_output=True,
                              env={**os.environ,'PYTHONIOENCODING':'cp1252','PYTHONUTF8':'0'},check=True)
        self.assertIn('٤٨',json.loads(result.stdout.decode('utf-8'))['blocks'][0]['text'])

    def test_docx_table_and_xlsx_sheet_references(self):
        from docx import Document
        d=Document();d.add_heading('Cargo',1);table=d.add_table(rows=2,cols=2);table.cell(0,0).text='Boxes';table.cell(1,0).text='48'
        stream=io.BytesIO();d.save(stream);blocks=extract('cargo.docx',stream.getvalue());self.assertTrue(any(b['table'] and b['reference']['heading']=='Cargo' for b in blocks))
        from openpyxl import Workbook
        w=Workbook();w.active.title='Cargo';w.active.append(['Pieces']);w.active.append([48]);stream=io.BytesIO();w.save(stream)
        blocks=extract('cargo.xlsx',stream.getvalue());self.assertEqual(blocks[1]['reference'],{'sheet':'Cargo','rowStart':2,'rowEnd':2})

    def test_text_pdf_page_reference(self):
        from pypdf import PdfWriter
        from pypdf.generic import DictionaryObject,NameObject,DecodedStreamObject
        writer=PdfWriter();page=writer.add_blank_page(400,400)
        font=DictionaryObject({NameObject('/Type'):NameObject('/Font'),NameObject('/Subtype'):NameObject('/Type1'),NameObject('/BaseFont'):NameObject('/Helvetica')})
        page[NameObject('/Resources')]=DictionaryObject({NameObject('/Font'):DictionaryObject({NameObject('/F1'):writer._add_object(font)})})
        content=DecodedStreamObject();content.set_data(b'BT /F1 12 Tf 10 350 Td (Shipment contains forty eight boxes in this test document.) Tj ET')
        page[NameObject('/Contents')]=writer._add_object(content);stream=io.BytesIO();writer.write(stream)
        blocks=extract('cargo.pdf',stream.getvalue());self.assertEqual(blocks[0]['reference']['page'],1);self.assertIn('eight boxes',blocks[0]['text'])

    def test_external_artifact_integrity_and_untracked_safety(self):
        store=root()
        with tempfile.TemporaryDirectory(dir=store/'eval') as directory:
            file=Path(directory)/'model.txt';file.write_text('safe')
            relative=file.relative_to(store).as_posix()
            with self.assertRaises(RuntimeError):reserve_destination(store,relative)
            register(store,file,'test','test','none','MIT');self.assertEqual(verify(store,relative),file)
            file.write_text('corrupt')
            with self.assertRaises(RuntimeError):verify(store,relative)
            from model_store import manifest,save_manifest
            data=manifest(store);data['files']=[f for f in data['files'] if f['path']!=relative];save_manifest(store,data)
        with self.assertRaises(ValueError):inside(store,'../outside')

    def test_hallucination_metric_rejects_invalid_json_and_fabricated_citation(self):
        row={'task':'grounded_qa','unanswerable':True,'target':{'statements':[{'label':'NOT FOUND','text':'not found','citations':[]}]},'schema':{'type':'object'},'context':[]}
        self.assertEqual(score(row,'not JSON')['hallucination'],1)
        self.assertEqual(score(row,json.dumps(row['target']))['not_found_correct'],1)
        self.assertEqual(score(row,'{"statements":["NOT FOUND"]}')['not_found_correct'],0)
        self.assertEqual(score(row,'{"statements":[{"label":"FACT","text":"invented","citations":["another-org"]}]}')['hallucination'],1)


if __name__=='__main__':unittest.main()
